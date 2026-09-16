from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from core.services import display_amount_from_ugx, get_user_currency


GREEN = "0F766E"
LIGHT = "E6FFFA"


def _style(ws, title, columns):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    ws.cell(1, 1, title).font = Font(size=14, bold=True, color="FFFFFF")
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=GREEN)
    ws.cell(1, 1).alignment = Alignment(horizontal="center")
    for col, value in enumerate(columns, 1):
        cell = ws.cell(3, col, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(columns))}3"


def _response(wb, filename):
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def _amount(value, request):
    return float(display_amount_from_ugx(value or 0, request=request))


def _currency_format(request):
    return f'"{get_user_currency(request)}" #,##0.00'


def budget_xlsx(budget, request=None):
    wb = Workbook(); ws = wb.active; ws.title = "Budget"
    columns = ["Code", "Outcome", "Activity", "Description", "Unit", "Price per unit", "Units", "Frequency", "Total", "Budget justification"]
    _style(ws, f"RICE West Nile - Budget: {budget.name}", columns)
    for row, line in enumerate(budget.lines.all(), 4):
        ws.append([line.code, line.outcome, line.activity, line.description, line.unit, _amount(line.price_per_unit, request), float(line.units), float(line.frequency), f"=F{row}*G{row}*H{row}", line.justification])
    for row in range(4, ws.max_row + 1):
        ws.cell(row, 6).number_format = ws.cell(row, 9).number_format = _currency_format(request)
    ws.cell(ws.max_row + 2, 8, "TOTAL").font = Font(bold=True)
    ws.cell(ws.max_row + 2, 9, f"=SUM(I4:I{ws.max_row})").font = Font(bold=True)
    for col in range(1, 11): ws.column_dimensions[get_column_letter(col)].width = min(38, max(13, max(len(str(ws.cell(r, col).value or "")) for r in range(3, ws.max_row + 1)) + 2))
    return _response(wb, f"budget_{budget.pk}.xlsx")


def performance_xlsx(budget, request=None):
    wb = Workbook(); ws = wb.active; ws.title = "Monthly Performance"
    columns = ["Code", "Activity", "Budget", "Funds received", "Expenditure", "Variance", "Donor balance", "Burn rate", "Absorption rate", "Comment"]
    _style(ws, f"RICE West Nile - Monthly Cumulative Budget Performance: {budget.name}", columns)
    for row, record in enumerate(budget.performance_records.select_related("line"), 4):
        ws.append([record.line.code, record.line.description, _amount(record.budget_amount, request), _amount(record.funds_received, request), _amount(record.expenditure, request), f"=C{row}-E{row}", f"=C{row}-D{row}", f'=IFERROR(E{row}/D{row},0)', f'=IFERROR(E{row}/C{row},0)', record.comment])
        for col in range(3, 8): ws.cell(row, col).number_format = _currency_format(request)
        ws.cell(row, 8).number_format = '0.0%'; ws.cell(row, 9).number_format = '0.0%'
    for col in range(1, 11): ws.column_dimensions[get_column_letter(col)].width = 18 if col not in (2, 10) else 32
    return _response(wb, f"budget_performance_{budget.pk}.xlsx")


def cashbook_xlsx(cash_book, request=None):
    wb = Workbook(); ws = wb.active; ws.title = "Cash book"
    columns = ["Date", "PV number", "Budget line", "Cheque number", "Payee", "Description of payment", "Receipt (Dr)", "Payment (Cr)", "Balance", "Comments"]
    _style(ws, f"RICE West Nile - Cash Book: {cash_book.name}", columns)
    ws.cell(4, 6, "Balance b/d"); ws.cell(4, 9, _amount(cash_book.opening_balance, request)); ws.cell(4, 9).number_format = _currency_format(request)
    for row, entry in enumerate(cash_book.entries.all(), 5):
        ws.append([entry.date, entry.pv_number, entry.budget_line, entry.cheque_number, entry.payee, entry.description, _amount(entry.receipt, request), _amount(entry.payment, request), f"=I{row-1}+G{row}-H{row}", entry.comments])
        ws.cell(row, 1).number_format = 'yyyy-mm-dd'; ws.cell(row, 7).number_format = ws.cell(row, 8).number_format = ws.cell(row, 9).number_format = _currency_format(request)
    for col in range(1, 11): ws.column_dimensions[get_column_letter(col)].width = 18 if col not in (6, 10) else 32
    return _response(wb, f"cash_book_{cash_book.pk}.xlsx")


def reconciliation_xlsx(reconciliation, request=None):
    wb = Workbook(); ws = wb.active; ws.title = "Bank reconciliation"
    columns = ["Date", "Type", "Reference", "Payee", "Amount", "Notes"]
    _style(ws, f"RICE West Nile - Bank Reconciliation: {reconciliation.cash_book.name}", columns)
    for row, item in enumerate(reconciliation.items.all(), 4):
        ws.append([item.date, item.get_item_type_display(), item.reference, item.payee, _amount(item.amount, request), item.notes])
        ws.cell(row, 1).number_format = 'yyyy-mm-dd'; ws.cell(row, 5).number_format = _currency_format(request)
    ws.cell(ws.max_row + 2, 1, "Statement balance"); ws.cell(ws.max_row, 5, _amount(reconciliation.statement_balance, request)); ws.cell(ws.max_row, 5).number_format = _currency_format(request)
    ws.cell(ws.max_row + 1, 1, "Adjusted balance"); ws.cell(ws.max_row, 5, _amount(reconciliation.adjusted_balance, request)); ws.cell(ws.max_row, 5).number_format = _currency_format(request)
    ws.cell(ws.max_row + 2, 1, "Difference"); ws.cell(ws.max_row, 5, _amount(reconciliation.difference, request)); ws.cell(ws.max_row, 5).number_format = _currency_format(request)
    return _response(wb, f"bank_reconciliation_{reconciliation.pk}.xlsx")
