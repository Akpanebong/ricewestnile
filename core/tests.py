from datetime import timedelta
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.utils import timezone
from django.urls import reverse
from docx import Document as WordDocument
from openpyxl import Workbook, load_workbook

from account.models import Department, Profile, ProgramArea, Unit
from .project_models import ActivityConceptNote, Project, ProjectActivity, ProjectDocument, ProjectSubmission, StakeholderUpdate, TeamMeeting
from .workplan_services import extract_workplan_activities


class ProgramAreaProjectDetailTests(TestCase):
    def setUp(self):
        self.user = Profile.objects.create_user(username="viewer", password="password")
        self.area = ProgramArea.objects.create(name="Test Program Area")
        self.project = Project.objects.create(name="Food Security", program_area=self.area)
        self.staff = Profile.objects.create_user(
            username="assigned-staff",
            first_name="Assigned",
            last_name="Staff",
            password="password",
            profile_type="Staff",
            project=self.project,
        )

    def test_program_area_detail_lists_projects_and_assigned_staff(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("program_area_detail", args=[self.area.pk]))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Food Security")
        self.assertContains(response, "Assigned Staff")

    def test_project_form_exposes_program_area(self):
        from .forms import ProjectForm

        self.assertIn("program_area", ProjectForm().fields)

    def test_unassigned_projects_are_available_to_employee_assignment(self):
        self.client.force_login(self.user)
        response = self.client.get("/units/projects/")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["projects"], [{"id": self.project.pk, "name": self.project.name}])

    def test_projects_can_be_filtered_by_program_area(self):
        other_area = ProgramArea.objects.create(name="Other Program Area")
        Project.objects.create(name="Other Project", program_area=other_area)
        self.client.force_login(self.user)

        response = self.client.get("/units/projects/", {"program_area": self.area.pk})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["projects"], [{"id": self.project.pk, "name": self.project.name}])

    def test_project_officer_can_upload_workplan(self):
        officer = Profile.objects.create_user(username="officer", password="password")
        self.project.project_officer = officer
        self.project.save()
        self.client.force_login(officer)

        response = self.client.post(
            reverse("core:project_document_upload", args=[self.project.pk, self.project.slug, "workplan"]),
            {"workplan_period": "monthly", "activities_count": 12, "file": SimpleUploadedFile("plan.pdf", b"pdf")},
        )

        document = ProjectDocument.objects.get(project=self.project)
        self.assertRedirects(response, reverse("core:project_activity_confirm", args=[self.project.pk, self.project.slug, document.pk]))
        self.assertEqual(document.document_type, ProjectDocument.WORKPLAN)
        self.assertIsNone(document.activities_count)

        response = self.client.post(
            reverse("core:project_activity_confirm", args=[self.project.pk, self.project.slug, document.pk]),
            {
                "activities-TOTAL_FORMS": "1",
                "activities-INITIAL_FORMS": "0",
                "activities-MIN_NUM_FORMS": "0",
                "activities-MAX_NUM_FORMS": "1000",
                "activities-0-name": "Community training",
                "activities-0-start_date": "2026-09-29",
                "activities-0-end_date": "2026-09-30",
            },
        )
        self.assertRedirects(response, reverse("core:project_documents", args=[self.project.pk, self.project.slug]))
        self.assertTrue(document.activities.filter(confirmed=True).exists())
        document.refresh_from_db()
        self.assertEqual(document.activities_count, 1)
        confirmed_activity = document.activities.get()
        self.assertTrue(confirmed_activity.confirmed)
        self.assertEqual(confirmed_activity.confirmed_by_id, officer.pk)
        self.assertIsNotNone(confirmed_activity.confirmed_at)

        from .forms import ProjectDocumentForm
        report_form = ProjectDocumentForm(document_type=ProjectDocument.REPORT, project=self.project)
        self.assertIn(confirmed_activity, report_form.fields["activity"].queryset)

    def test_non_officer_cannot_upload_project_document(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("core:project_document_upload", args=[self.project.pk, self.project.slug, "report"]),
            {"report_type": "monthly_report", "file": SimpleUploadedFile("report.pdf", b"pdf")},
        )

        self.assertRedirects(response, reverse("core:project_documents", args=[self.project.pk, self.project.slug]))
        self.assertFalse(ProjectDocument.objects.filter(project=self.project).exists())

    def test_activity_status_uses_dates_and_linked_activity_report(self):
        workplan = ProjectDocument.objects.create(
            project=self.project,
            document_type=ProjectDocument.WORKPLAN,
            workplan_period=ProjectDocument.MONTHLY,
            file=SimpleUploadedFile("plan.pdf", b"pdf"),
        )
        activity = ProjectActivity.objects.create(
            workplan=workplan,
            name="Distribute seeds",
            start_date=timezone.localdate(),
            end_date=timezone.localdate(),
            confirmed=True,
        )
        self.assertEqual(activity.status, ProjectActivity.IN_PROGRESS)

        report = ProjectDocument.objects.create(
            project=self.project,
            document_type=ProjectDocument.REPORT,
            report_type=ProjectDocument.ACTIVITY_REPORT,
            activity=activity,
            file=SimpleUploadedFile("activity-1.pdf", b"pdf"),
        )
        report.uploaded_at = timezone.now()
        report.save(update_fields=["uploaded_at"])
        self.assertEqual(activity.status, ProjectActivity.ON_SCHEDULE)

        activity.start_date = timezone.localdate() - timedelta(days=3)
        activity.end_date = timezone.localdate() - timedelta(days=1)
        activity.save(update_fields=["start_date", "end_date"])
        report.uploaded_at = timezone.now()
        report.save(update_fields=["uploaded_at"])
        self.assertEqual(activity.status, ProjectActivity.OFF_SCHEDULE)
        self.assertEqual(self.project.schedule_status, "Off Schedule")

    def test_excel_workplan_activities_are_extracted_as_drafts(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["Activity Name", "Start Date", "End Date"])
        sheet.append(["Community training", timezone.localdate(), timezone.localdate()])
        stream = BytesIO()
        workbook.save(stream)
        workplan = ProjectDocument.objects.create(
            project=self.project,
            document_type=ProjectDocument.WORKPLAN,
            workplan_period=ProjectDocument.MONTHLY,
            file=SimpleUploadedFile("plan.xlsx", stream.getvalue()),
        )

        self.assertEqual(extract_workplan_activities(workplan), 1)
        activity = workplan.activities.get()
        self.assertEqual(activity.name, "Community training")
        self.assertFalse(activity.confirmed)

    def test_docx_workplan_activities_are_extracted(self):
        word_document = WordDocument()
        word_document.add_paragraph("Workplan for October")
        word_document.add_paragraph("Activity 1: Community training")
        word_document.add_paragraph("Start date: 12/10/2026\t\tEnd date: 17/10/2026")
        word_document.add_paragraph("Activity 2: Review exam questions")
        word_document.add_paragraph("Start date: 20/10/2026\t\tEnd date: 25/10/2026")
        stream = BytesIO()
        word_document.save(stream)
        workplan = ProjectDocument.objects.create(
            project=self.project,
            document_type=ProjectDocument.WORKPLAN,
            workplan_period=ProjectDocument.MONTHLY,
            file=SimpleUploadedFile("workplan.docx", stream.getvalue()),
        )

        self.assertEqual(extract_workplan_activities(workplan), 2)
        self.assertEqual(workplan.activities.first().start_date.isoformat(), "2026-10-12")

    def test_docx_table_variants_extract_code_activity_and_description(self):
        word_document = WordDocument()
        table = word_document.add_table(rows=1, cols=5)
        for cell, value in zip(table.rows[0].cells, ["Activity Code", "Activities", "Activity Description", "Start Date", "End Date"]):
            cell.text = value
        row = table.add_row().cells
        for cell, value in zip(row, ["ACT-01", "Community training", "Train farmer groups", "12/10/2026", "17/10/2026"]):
            cell.text = value
        stream = BytesIO()
        word_document.save(stream)
        workplan = ProjectDocument.objects.create(
            project=self.project,
            document_type=ProjectDocument.WORKPLAN,
            workplan_period=ProjectDocument.MONTHLY,
            file=SimpleUploadedFile("table-workplan.docx", stream.getvalue()),
        )

        self.assertEqual(extract_workplan_activities(workplan), 1)
        activity = workplan.activities.get()
        self.assertEqual(activity.code, "ACT-01")
        self.assertEqual(activity.name, "Community training")

    def test_report_activity_choices_are_limited_to_the_same_project(self):
        from .forms import ProjectDocumentForm

        other_project = Project.objects.create(name="Other Project")
        other_workplan = ProjectDocument.objects.create(
            project=other_project,
            document_type=ProjectDocument.WORKPLAN,
            workplan_period=ProjectDocument.MONTHLY,
            file=SimpleUploadedFile("other.pdf", b"pdf"),
        )
        other_activity = ProjectActivity.objects.create(
            workplan=other_workplan,
            name="Other activity",
            start_date=timezone.localdate(),
            end_date=timezone.localdate(),
            confirmed=True,
        )
        form = ProjectDocumentForm(document_type=ProjectDocument.REPORT, project=self.project)
        self.assertNotIn(other_activity, form.fields["activity"].queryset)

    def test_project_without_workplan_is_not_started(self):
        self.assertEqual(self.project.schedule_status, "Not Started")

    def test_project_officer_can_enter_project_reporting_records(self):
        officer = Profile.objects.create_user(username="reporting-officer", password="password")
        self.project.project_officer = officer
        self.project.save()
        self.client.force_login(officer)

        response = self.client.post(
            reverse("core:project_reporting_input", args=[self.project.pk, self.project.slug]),
            {
                "form_type": "stakeholder",
                "stakeholder-date": "2026-09-29",
                "stakeholder-stakeholder_name": "District Local Government",
                "stakeholder-category": "government",
                "stakeholder-engagement_method": "meeting",
                "stakeholder-update_summary": "Reviewed implementation progress.",
                "stakeholder-follow_up_status": "ongoing",
            },
        )
        self.assertRedirects(response, reverse("core:project_reporting_input", args=[self.project.pk, self.project.slug]))
        self.assertTrue(StakeholderUpdate.objects.filter(project=self.project).exists())

        response = self.client.post(
            reverse("core:project_reporting_input", args=[self.project.pk, self.project.slug]),
            {
                "form_type": "meeting",
                "meeting-meeting_date": "2026-09-29",
                "meeting-start_time": "09:00",
                "meeting-end_time": "10:00",
                "meeting-key_action_items": "Submit the monthly narrative.",
                "meeting-status": "pending",
            },
        )
        self.assertRedirects(response, reverse("core:project_reporting_input", args=[self.project.pk, self.project.slug]))
        self.assertTrue(TeamMeeting.objects.filter(project=self.project).exists())

    def test_submission_file_types_and_cmt_only_view(self):
        officer = Profile.objects.create_user(username="submission-officer", password="password")
        self.project.project_officer = officer
        self.project.save()
        self.client.force_login(officer)
        url = reverse("core:project_reporting_input", args=[self.project.pk, self.project.slug])

        invalid = self.client.post(
            url,
            {
                "form_type": "submission",
                "submission-document_type": ProjectSubmission.NARRATIVE,
                "submission-reporting_period": "September 2026",
                "submission-file": SimpleUploadedFile("narrative.xlsx", b"not a pdf"),
            },
        )
        self.assertEqual(invalid.status_code, 200)
        self.assertFalse(ProjectSubmission.objects.exists())

        valid = self.client.post(
            url,
            {
                "form_type": "submission",
                "submission-document_type": ProjectSubmission.NARRATIVE,
                "submission-reporting_period": "September 2026",
                "submission-file": SimpleUploadedFile("narrative.pdf", b"pdf content"),
            },
        )
        self.assertRedirects(valid, url)
        submission = ProjectSubmission.objects.get(project=self.project)

        self.assertEqual(self.client.get(reverse("core:cmt_project_reporting", args=[self.project.pk, self.project.slug])).status_code, 403)

        cmt = Profile.objects.create_user(username="cmt", password="password", is_CMT=True)
        self.client.force_login(cmt)
        self.assertEqual(self.client.get(reverse("core:cmt_project_reporting", args=[self.project.pk, self.project.slug])).status_code, 200)
        self.assertContains(
            self.client.get(reverse("core:cmt_project_reporting", args=[self.project.pk, self.project.slug]) + "?section=documents"),
            "Project Documents",
        )
        self.assertNotContains(
            self.client.get(reverse("core:cmt_project_reporting", args=[self.project.pk, self.project.slug]) + "?section=documents"),
            "Stakeholder Updates</h5>",
        )
        download = self.client.get(reverse("core:cmt_project_submission_download", args=[self.project.pk, self.project.slug, submission.pk]))
        self.assertEqual(download.status_code, 200)

    def test_cmt_can_download_tracker_records_in_template_formats(self):
        cmt = Profile.objects.create_user(username="export-cmt", password="password", is_CMT=True)
        StakeholderUpdate.objects.create(
            project=self.project, date=timezone.localdate(), stakeholder_name="District", category="government",
            engagement_method="meeting", update_summary="Progress update", entered_by=cmt,
        )
        TeamMeeting.objects.create(
            project=self.project, meeting_date=timezone.localdate(), start_time="09:00", end_time="10:00",
            key_action_items="Follow up", entered_by=cmt,
        )
        self.client.force_login(cmt)

        stakeholder_response = self.client.get(reverse("core:cmt_stakeholder_export", args=[self.project.pk, self.project.slug]))
        self.assertEqual(stakeholder_response.status_code, 200)
        self.assertIn("stakeholder-updates.xlsx", stakeholder_response["Content-Disposition"])
        stakeholder_book = load_workbook(BytesIO(stakeholder_response.content), data_only=True)
        self.assertEqual(stakeholder_book["Sheet1"].cell(2, 2).value, "District")

        meeting_response = self.client.get(reverse("core:cmt_team_meetings_export", args=[self.project.pk, self.project.slug]))
        self.assertEqual(meeting_response.status_code, 200)
        self.assertIn("team-meeting-tracker.xlsx", meeting_response["Content-Disposition"])
        meeting_book = load_workbook(BytesIO(meeting_response.content), data_only=True)
        self.assertEqual(meeting_book["Sept"].cell(4, 4).value, self.project.name)
        self.assertEqual(meeting_book["Sept"].cell(4, 2).number_format, "h:mm AM/PM")

        all_stakeholder = self.client.get(reverse("core:cmt_all_stakeholder_export"))
        self.assertEqual(all_stakeholder.status_code, 200)
        self.assertIn("all-stakeholder-updates.xlsx", all_stakeholder["Content-Disposition"])
        all_meetings = self.client.get(reverse("core:cmt_all_team_meetings_export"))
        self.assertEqual(all_meetings.status_code, 200)
        self.assertIn("all-team-meetings.xlsx", all_meetings["Content-Disposition"])


class ActivityConceptNoteWorkflowTests(TestCase):
    def setUp(self):
        self.officer = Profile.objects.create_user(username="concept-po", password="password")
        self.project_head = Profile.objects.create_user(username="concept-project-head", password="password")
        self.program_head = Profile.objects.create_user(username="concept-program-head", password="password")
        department = Department.objects.create(name="Programs", head=self.program_head)
        unit = Unit.objects.create(department=department, name="Program Unit")
        area = ProgramArea.objects.create(name="Concept Program")
        self.project = Project.objects.create(
            name="Concept Note Project", program_area=area, unit=unit,
            project_officer=self.officer, project_head=self.project_head,
        )
        workplan = ProjectDocument.objects.create(
            project=self.project, document_type=ProjectDocument.WORKPLAN,
            workplan_period=ProjectDocument.MONTHLY,
            file=SimpleUploadedFile("plan.pdf", b"pdf"),
        )
        self.activity = ProjectActivity.objects.create(
            workplan=workplan, code="ACT-01", name="Community dialogue",
            start_date=timezone.localdate(), end_date=timezone.localdate(), confirmed=True,
        )

    def test_po_online_note_moves_through_review_and_approval(self):
        self.client.force_login(self.officer)
        detail_url = reverse("core:activity_concept_note_detail", args=[self.project.pk, self.project.slug, self.activity.pk])
        online_url = reverse("core:activity_concept_note_entry", args=[self.project.pk, self.project.slug, self.activity.pk])
        online_page = self.client.get(online_url)
        self.assertContains(online_page, "A. Profile")
        self.assertContains(online_page, "B. Description")
        self.assertContains(online_page, "form-control")
        response = self.client.post(online_url, {
            "project_name": self.project.name, "activity_code": "ACT-01", "activity_name": "Community dialogue",
            "core_program_area": "Concept Program", "outcome_result_area": "Improved participation",
            "activity_start_date": timezone.localdate(), "activity_end_date": timezone.localdate(),
            "concept_date": timezone.localdate(), "background": "A documented need.",
            "objectives": "Improve participation.", "methodology": "Dialogue.",
            "expected_outputs": "One dialogue completed.", "inputs_required": "Facilitators.",
            "tasks_text": "Mobilise participants | 2026-10-01 | Project Officer",
            "budget_lines_text": "Refreshments | lot | 1 | 1 | 100000 | 100000",
            "total_amount_in_words": "One hundred thousand shillings",
        })
        self.assertRedirects(response, detail_url)
        note = ActivityConceptNote.objects.get(activity=self.activity)
        self.assertEqual(note.status, ActivityConceptNote.DRAFT)

        self.client.post(reverse("core:activity_concept_note_submit", args=[self.project.pk, self.project.slug, self.activity.pk]))
        note.refresh_from_db()
        self.assertEqual(note.status, ActivityConceptNote.SUBMITTED)

        self.client.force_login(self.project_head)
        self.client.post(reverse("core:activity_concept_note_review", args=[self.project.pk, self.project.slug, self.activity.pk]), {"action": "review", "comments": "Looks good."})
        note.refresh_from_db()
        self.assertEqual(note.status, ActivityConceptNote.REVIEWED)

        self.client.force_login(self.program_head)
        self.client.post(reverse("core:activity_concept_note_approve", args=[self.project.pk, self.project.slug, self.activity.pk]))
        note.refresh_from_db()
        self.assertEqual(note.status, ActivityConceptNote.APPROVED)

    def test_po_can_upload_completed_form(self):
        self.client.force_login(self.officer)
        upload_url = reverse("core:activity_concept_note_upload", args=[self.project.pk, self.project.slug, self.activity.pk])
        upload_page = self.client.get(upload_url)
        self.assertEqual(upload_page.status_code, 200)
        self.assertContains(upload_page, "Upload completed concept note")
        self.assertContains(upload_page, reverse("core:project_documents", args=[self.project.pk, self.project.slug]))
        response = self.client.post(reverse("core:activity_concept_note_upload", args=[self.project.pk, self.project.slug, self.activity.pk]), {
            "uploaded_file": SimpleUploadedFile("concept-note.docx", b"docx content"),
        })
        self.assertEqual(response.status_code, 302)
        note = ActivityConceptNote.objects.get(activity=self.activity)
        self.assertEqual(note.source, ActivityConceptNote.UPLOADED)
        self.assertIn("concept-note", note.uploaded_file.name)
