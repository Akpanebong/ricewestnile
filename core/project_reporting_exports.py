from copy import copy
from io import BytesIO
from pathlib import Path

from django.conf import settings
from django.http import HttpResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .project_models import StakeholderUpdate, TeamMeeting


def _template_path(filename):
    return Path(settings.BASE_DIR) / "core" / filename


def _copy_row_style(source_row, target_row):
    for source_cell, target_cell in zip(source_row, target_row):
        if source_cell.has_style:
            target_cell._style = copy(source_cell._style)
        target_cell.number_format = source_cell.number_format
        target_cell.alignment = copy(source_cell.alignment)
        target_cell.protection = copy(source_cell.protection)


def _styled_header(sheet, row_number, headers):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for column, value in enumerate(headers, start=1):
        cell = sheet.cell(row_number, column, value)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[row_number].height = 32


def _load_stakeholder_template():
    path = _template_path("Stakeholder_Updates_Tracking_Template.xlsx")
    if path.exists():
        return load_workbook(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sheet1"
    _styled_header(sheet, 1, [
        "Date", "Stakeholder Name", "Category (Govt/Donor/Community/CSO/Private Sector etc.)",
        "Designation", "Email/Phone", "Engagement Method (Meeting/Call/Email etc.)",
        "Update / Discussion Summary", "Decisions / Action Points", "Responsible Person ",
        "Deadline", "Follow-up Status (Pending/Ongoing/Completed)", "Remarks/Notes",
    ])
    categories = workbook.create_sheet("Sheet2")
    _styled_header(categories, 1, ["Category name"])
    for row, value in enumerate(["Government", "Donor", "Community", "CSO", "Private Sector"], start=2):
        categories.cell(row, 1, value)
    return workbook


def _load_team_template():
    path = _template_path("Team Meeting Tracker.xlsx")
    if path.exists():
        return load_workbook(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Sept"
    headers = [
        "Meeting Date", "Start Time", "End Time", "Team/ Project Name", "Chairperson",
        "Minute Taker", "Participants (M/F)", "Key Action Items", "Timeline", "Status",
        "Area of support", "Person Responsible", "",
    ]
    for month in ("Sept", "Oct", "Nov", "Dec"):
        sheet = workbook[month] if month == "Sept" else workbook.create_sheet(month)
        sheet["A1"] = "Team Meeting Tracker "
        sheet["H2"] = "Summary of the Action Plan"
        sheet["K2"] = "Support Required"
        _styled_header(sheet, 3, headers)
    reference = workbook.create_sheet("Sheet2")
    _styled_header(reference, 1, ["Project Name"])
    return workbook


def stakeholder_updates_workbook(project):
    workbook = _load_stakeholder_template()
    sheet = workbook["Sheet1"]
    records = StakeholderUpdate.objects.select_related("project").order_by("date", "created_at")
    if project is not None:
        records = records.filter(project=project)
    first_data_row = 2
    template_row = list(sheet[first_data_row])
    for index, record in enumerate(records, start=first_data_row):
        if index > sheet.max_row:
            sheet.append([None] * 12)
            _copy_row_style(template_row, list(sheet[index]))
        row_values = [
            record.date, record.stakeholder_name, record.get_category_display(), record.designation,
            record.contact, record.get_engagement_method_display(), record.update_summary,
            record.action_points, record.responsible_person, record.deadline,
            record.get_follow_up_status_display(), record.remarks,
        ]
        for column, value in enumerate(row_values, start=1):
            sheet.cell(index, column).value = value
            if column in {1, 10}:
                sheet.cell(index, column).number_format = "dd mmm yyyy"
    for column, width in {"A": 14, "B": 24, "C": 24, "D": 20, "E": 22, "F": 24, "G": 36, "H": 36, "I": 22, "J": 14, "K": 24, "L": 30}.items():
        sheet.column_dimensions[column].width = width
    return _as_xlsx_response(workbook, "all-stakeholder-updates.xlsx" if project is None else "stakeholder-updates.xlsx")


def team_meetings_workbook(project):
    workbook = _load_team_template()
    records = TeamMeeting.objects.select_related("project").order_by("meeting_date", "start_time")
    if project is not None:
        records = records.filter(project=project)
    month_sheets = {"Jan": "Jan", "Feb": "Feb", "Mar": "Mar", "Apr": "Apr", "May": "May", "Jun": "Jun", "Jul": "Jul", "Aug": "Aug", "Sep": "Sept", "Oct": "Oct", "Nov": "Nov", "Dec": "Dec"}
    template_sheet = workbook["Sept"]
    rows_by_sheet = {}
    for record in records:
        sheet_name = month_sheets[record.meeting_date.strftime("%b")]
        if sheet_name not in workbook.sheetnames:
            sheet = workbook.copy_worksheet(template_sheet)
            sheet.title = sheet_name
        rows_by_sheet.setdefault(sheet_name, []).append(record)

    for sheet_name, sheet_records in rows_by_sheet.items():
        sheet = workbook[sheet_name]
        template_row = list(sheet[4])
        for row_number, record in enumerate(sheet_records, start=4):
            if row_number > sheet.max_row:
                sheet.append([None] * 13)
                _copy_row_style(template_row, list(sheet[row_number]))
            row_values = [
                record.meeting_date, record.start_time, record.end_time, record.project.name,
                record.chairperson, record.minute_taker, record.participants_m_f,
                record.key_action_items, record.timeline, record.get_status_display(),
                record.area_of_support, record.person_responsible, None,
            ]
            for column, value in enumerate(row_values, start=1):
                sheet.cell(row_number, column).value = value
            sheet.cell(row_number, 1).number_format = "dd mmm yyyy"
            sheet.cell(row_number, 2).number_format = "h:mm AM/PM"
            sheet.cell(row_number, 3).number_format = "h:mm AM/PM"
        for column, width in {"A": 16, "B": 14, "C": 14, "D": 28, "E": 22, "F": 22, "G": 18, "H": 36, "I": 22, "J": 16, "K": 30, "L": 24}.items():
            sheet.column_dimensions[column].width = width
    return _as_xlsx_response(workbook, "all-team-meetings.xlsx" if project is None else "team-meeting-tracker.xlsx")


def _as_xlsx_response(workbook, filename):
    stream = BytesIO()
    workbook.save(stream)
    response = HttpResponse(
        stream.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
